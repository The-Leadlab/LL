"""
Fetch Serendien client leads from LeadLab production, run internal psychometrics,
optionally persist to each lead, and export Excel with pitch personalization.

Usage (from backend/):
  py scripts/serendien_psychometrics_export.py
"""
from __future__ import annotations

import json
import math
import os
import sys
import time
from collections import Counter
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Dict, List, Optional

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

from app.services.internal_psychometric_analyzer import InternalPsychometricAnalyzer

API = os.environ.get("LEADLAB_API", "https://api.the-leadlab.com/api/v1")
EMAIL = os.environ.get("LEADLAB_EMAIL", "ali@the-leadlab.com")
PASSWORD = os.environ.get("LEADLAB_PASSWORD", "")
CLIENT_NAME = os.environ.get("LEADLAB_CLIENT", "Serendien")
OUT = Path(os.environ.get(
    "LEADLAB_PSYCHO_OUT",
    r"c:\Users\attia\Downloads\SERENDIEN_Leads_Psychometrics_Pitch.xlsx",
))
PERSIST = os.environ.get("LEADLAB_PERSIST", "1") != "0"

PITCH_BY_TYPE = {
    "D": {
        "pitch_style": "Results-first, concise, ROI-led",
        "opener_angle": "Lead with outcomes, time saved, and competitive edge — skip small talk.",
        "do": "Be direct; quantify impact; offer clear next step",
        "avoid": "Long intros, soft hedging, vague timelines",
    },
    "I": {
        "pitch_style": "Energetic, relationship and vision-led",
        "opener_angle": "Open with a story or social proof; make the opportunity feel exciting.",
        "do": "Use enthusiasm, testimonials, collaboration language",
        "avoid": "Dry data dumps, overly formal tone, ignoring rapport",
    },
    "S": {
        "pitch_style": "Steady, trust-building, low-pressure",
        "opener_angle": "Emphasize reliability, continuity, and support — reduce perceived risk.",
        "do": "Go slow; show process; offer reassurance and references",
        "avoid": "Hard closes, surprise changes, aggressive urgency",
    },
    "C": {
        "pitch_style": "Evidence-led, precise, structured",
        "opener_angle": "Lead with proof, methodology, and clear specs — invite scrutiny.",
        "do": "Share data, comparisons, case studies with detail",
        "avoid": "Hype, fluff, incomplete answers, pressure to decide fast",
    },
}


def pitch_for(personality_type: str) -> dict:
    primary = (personality_type or "S")[0].upper()
    if primary not in PITCH_BY_TYPE:
        primary = "S"
    return PITCH_BY_TYPE[primary]


def s(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and math.isnan(v):
        return ""
    return str(v).strip()


def login() -> str:
    if not PASSWORD:
        raise RuntimeError("Set LEADLAB_PASSWORD in the environment before running this script.")
    r = requests.post(
        f"{API}/auth/login",
        json={"email": EMAIL, "password": PASSWORD},
        timeout=60,
    )
    r.raise_for_status()
    data = r.json()
    token = data.get("access_token") or data.get("token")
    if not token:
        raise RuntimeError(f"No token in login response: {data}")
    return token


def find_client_id(session: requests.Session) -> int:
    r = session.get(f"{API}/clients/", params={"include_archived": False}, timeout=60)
    r.raise_for_status()
    items = r.json().get("items") or []
    for c in items:
        if (c.get("name") or "").strip().lower() == CLIENT_NAME.lower():
            return int(c["id"])
    names = [c.get("name") for c in items]
    raise RuntimeError(f"Client '{CLIENT_NAME}' not found. Have: {names}")


def fetch_all_leads(session: requests.Session, client_id: int) -> List[dict]:
    leads: List[dict] = []
    skip = 0
    limit = 200
    while True:
        r = session.get(
            f"{API}/leads/",
            params={
                "client_id": client_id,
                "skip": skip,
                "limit": limit,
                "sort_by": "created_at",
                "sort_desc": True,
            },
            timeout=120,
        )
        r.raise_for_status()
        payload = r.json()
        batch = payload.get("results") or []
        leads.extend(batch)
        if not batch or not payload.get("has_more"):
            break
        skip += len(batch)
    return leads


def analyze_one(analyzer: InternalPsychometricAnalyzer, lead: dict) -> Dict[str, Any]:
    lead_data = {
        "job_title": s(lead.get("job_title")),
        "company": s(lead.get("company")),
        "industry": s(lead.get("sector")),
        "first_name": s(lead.get("first_name")),
        "last_name": s(lead.get("last_name")),
        "email": s(lead.get("email")),
        "linkedin_url": s(lead.get("linkedin")),
        "enhanced_linkedin": {},
    }
    disc = analyzer._analyze_disc_profile(lead_data)
    big5 = analyzer._analyze_big_five(lead_data)
    personality_type = analyzer._determine_personality_type(disc, big5)
    communication = analyzer._analyze_communication_style(lead_data, personality_type)
    sales = analyzer._generate_sales_intelligence(personality_type, lead_data)
    behavioral = analyzer._generate_behavioral_predictions(personality_type)
    confidence = analyzer._calculate_confidence_score(lead_data)
    description = analyzer._get_personality_description(personality_type)
    strengths = analyzer._identify_strengths(personality_type, disc)
    pitch = pitch_for(personality_type)
    primary = (personality_type or "S")[0].upper()

    psycho_payload = {
        "personality_type": personality_type,
        "primary_disc": primary,
        "disc": {k: round(disc.get(k, 0), 1) for k in ("D", "I", "S", "C")},
        "big_five": {k: round(big5.get(k, 0), 2) for k in (
            "openness", "conscientiousness", "extraversion", "agreeableness", "neuroticism"
        )},
        "communication_style": communication.get("primary_style"),
        "sales_approach": sales.get("primary_approach"),
        "confidence_score": confidence,
        "source": "internal_analyzer_serendien_export",
    }

    return {
        "lead_id": lead.get("id"),
        "first_name": lead.get("first_name"),
        "last_name": lead.get("last_name"),
        "full_name": lead.get("full_name") or f"{lead.get('first_name') or ''} {lead.get('last_name') or ''}".strip(),
        "company": lead.get("company"),
        "job_title": lead.get("job_title"),
        "email": lead.get("email"),
        "telephone": lead.get("telephone") or lead.get("mobile"),
        "linkedin": lead.get("linkedin"),
        "sector": lead.get("sector"),
        "country": lead.get("country"),
        "location": lead.get("location"),
        "personality_type": personality_type,
        "primary_disc": primary,
        "personality_description": description,
        "disc_D": round(disc.get("D", 0), 1),
        "disc_I": round(disc.get("I", 0), 1),
        "disc_S": round(disc.get("S", 0), 1),
        "disc_C": round(disc.get("C", 0), 1),
        "big5_openness": round(big5.get("openness", 0), 2),
        "big5_conscientiousness": round(big5.get("conscientiousness", 0), 2),
        "big5_extraversion": round(big5.get("extraversion", 0), 2),
        "big5_agreeableness": round(big5.get("agreeableness", 0), 2),
        "big5_neuroticism": round(big5.get("neuroticism", 0), 2),
        "communication_style": communication.get("primary_style"),
        "sales_approach": sales.get("primary_approach"),
        "decision_style": sales.get("decision_making_style"),
        "motivators": "; ".join(sales.get("key_motivators") or [])
        if isinstance(sales.get("key_motivators"), list)
        else sales.get("key_motivators"),
        "likely_objections": "; ".join(sales.get("likely_objections") or [])
        if isinstance(sales.get("likely_objections"), list)
        else sales.get("likely_objections"),
        "strengths": "; ".join(strengths) if isinstance(strengths, list) else strengths,
        "work_style": (behavioral or {}).get("work_style"),
        "confidence_score": round(confidence, 2) if isinstance(confidence, (int, float)) else confidence,
        "pitch_style": pitch["pitch_style"],
        "pitch_opener_angle": pitch["opener_angle"],
        "pitch_do": pitch["do"],
        "pitch_avoid": pitch["avoid"],
        "_psycho_payload": psycho_payload,
    }


def persist_psychometrics(session: requests.Session, lead_id: int, payload: dict) -> bool:
    try:
        r = session.put(
            f"{API}/leads/{lead_id}",
            json={"psychometrics": payload},
            timeout=60,
        )
        return r.status_code < 400
    except Exception:
        return False


def write_excel(rows: List[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Psychometrics + Pitch"
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")

    keys = [k for k in rows[0].keys() if not k.startswith("_")] if rows else []
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append([r.get(k) for k in keys])
    for col in range(1, len(keys) + 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = min(32, max(12, len(keys[col - 1]) + 2))

    summary = wb.create_sheet("Profile Summary", 0)
    c = Counter((r.get("primary_disc") or "?") for r in rows)
    summary.append(["Primary DISC", "Count", "Share %", "Pitch style"])
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = header_font
    total = len(rows) or 1
    for disc, n in sorted(c.items(), key=lambda x: -x[1]):
        pitch = PITCH_BY_TYPE.get(disc, {}).get("pitch_style", "")
        summary.append([disc, n, round(100.0 * n / total, 1), pitch])
    summary.append([])
    summary.append(["Total leads", len(rows)])
    summary.append(["Client", CLIENT_NAME])
    summary.column_dimensions["A"].width = 16
    summary.column_dimensions["B"].width = 10
    summary.column_dimensions["C"].width = 12
    summary.column_dimensions["D"].width = 40

    playbook = wb.create_sheet("Pitch Playbook by DISC")
    playbook.append(["DISC", "Pitch style", "Opener angle", "Do", "Avoid"])
    for cell in playbook[1]:
        cell.fill = header_fill
        cell.font = header_font
    for disc, p in PITCH_BY_TYPE.items():
        playbook.append([disc, p["pitch_style"], p["opener_angle"], p["do"], p["avoid"]])
    for col in range(1, 6):
        playbook.column_dimensions[get_column_letter(col)].width = 36

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> int:
    print(f"Logging in as {EMAIL}…")
    token = login()
    session = requests.Session()
    session.headers.update({"Authorization": f"Bearer {token}"})

    client_id = find_client_id(session)
    print(f"Client '{CLIENT_NAME}' id={client_id}")
    leads = fetch_all_leads(session, client_id)
    print(f"Fetched {len(leads)} leads")
    if not leads:
        print("No leads — nothing to do")
        return 1

    analyzer = InternalPsychometricAnalyzer()
    rows: List[dict] = []
    persisted = 0
    for i, lead in enumerate(leads, 1):
        row = analyze_one(analyzer, lead)
        if PERSIST:
            ok = persist_psychometrics(session, int(lead["id"]), row["_psycho_payload"])
            if ok:
                persisted += 1
        rows.append(row)
        if i % 50 == 0 or i == len(leads):
            print(f"Analyzed {i}/{len(leads)} (persisted={persisted})")

    write_excel(rows, OUT)
    print(f"Wrote {OUT}")
    print(f"Folder: {OUT.parent}")
    counts = Counter(r.get("primary_disc") for r in rows)
    print("DISC mix:", dict(counts))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
