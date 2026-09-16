"""
Offline psychometrics for Serenidien prospects using LeadLab InternalPsychometricAnalyzer.
Does not require API login or DB. Writes Excel to Downloads.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path
from types import SimpleNamespace

# Ensure backend package imports work
BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from app.services.internal_psychometric_analyzer import InternalPsychometricAnalyzer


PITCH_BY_TYPE = {
    "D": {
        "pitch_style": "Results-first, concise, ROI-led",
        "opener_angle": "Lead with outcomes, time saved, and competitive edge — skip small talk.",
        "do": "Be direct; quantify impact; offer clear next step",
        "avoid": "Long intros, soft hedging, vague timelines",
    },
    "I": {
        "pitch_style": "Energetic, relationship and vision-led",
        "opener_angle": "Open with a story or social proof; make the opportunity feel exciting.",
        "do": "Use enthusiasm, testimonials, collaboration language",
        "avoid": "Dry data dumps, overly formal tone, ignoring rapport",
    },
    "S": {
        "pitch_style": "Steady, trust-building, low-pressure",
        "opener_angle": "Emphasize reliability, continuity, and support — reduce perceived risk.",
        "do": "Go slow; show process; offer reassurance and references",
        "avoid": "Hard closes, surprise changes, aggressive urgency",
    },
    "C": {
        "pitch_style": "Evidence-led, precise, structured",
        "opener_angle": "Lead with proof, methodology, and clear specs — invite scrutiny.",
        "do": "Share data, comparisons, case studies with detail",
        "avoid": "Hype, fluff, incomplete answers, pressure to decide fast",
    },
}


def pitch_for(personality_type: str) -> dict:
    primary = (personality_type or "S")[0].upper()
    if primary not in PITCH_BY_TYPE:
        primary = "S"
    return PITCH_BY_TYPE[primary]


def analyze_one(analyzer: InternalPsychometricAnalyzer, lead: dict, idx: int) -> dict:
    mock = SimpleNamespace(
        id=idx,
        first_name=lead.get("first_name"),
        last_name=lead.get("last_name"),
        email=lead.get("email"),
        job_title=lead.get("job_title"),
        company=lead.get("company"),
        linkedin=lead.get("linkedin"),
        sector=lead.get("sector"),
        industry=lead.get("sector"),
        website=lead.get("website"),
        country=lead.get("country"),
        location=lead.get("location"),
        telephone=lead.get("telephone"),
        mobile=lead.get("mobile"),
        organization_id=None,
    )

    def s(v) -> str:
        if v is None:
            return ""
        if isinstance(v, float):
            # JSON/Excel NaN
            import math

            if math.isnan(v):
                return ""
        return str(v).strip()

    lead_data = {
        "job_title": s(lead.get("job_title")),
        "company": s(lead.get("company")),
        "industry": s(lead.get("sector")),
        "first_name": s(lead.get("first_name")),
        "last_name": s(lead.get("last_name")),
        "email": s(lead.get("email")),
        "linkedin_url": s(lead.get("linkedin")),
        "enhanced_linkedin": {},
    }

    disc = analyzer._analyze_disc_profile(lead_data)
    big5 = analyzer._analyze_big_five(lead_data)
    personality_type = analyzer._determine_personality_type(disc, big5)
    communication = analyzer._analyze_communication_style(lead_data, personality_type)
    sales = analyzer._generate_sales_intelligence(personality_type, lead_data)
    behavioral = analyzer._generate_behavioral_predictions(personality_type)
    confidence = analyzer._calculate_confidence_score(lead_data)
    description = analyzer._get_personality_description(personality_type)
    strengths = analyzer._identify_strengths(personality_type, disc)
    pitch = pitch_for(personality_type)

    return {
        "first_name": lead.get("first_name"),
        "last_name": lead.get("last_name"),
        "company": lead.get("company"),
        "job_title": lead.get("job_title"),
        "email": lead.get("email"),
        "telephone": lead.get("telephone"),
        "linkedin": lead.get("linkedin"),
        "sector": lead.get("sector"),
        "country": lead.get("country"),
        "city": lead.get("location"),
        "source_persona_sheet": lead.get("source_persona_sheet"),
        "personality_type": personality_type,
        "personality_description": description,
        "disc_D": round(disc.get("D", 0), 1),
        "disc_I": round(disc.get("I", 0), 1),
        "disc_S": round(disc.get("S", 0), 1),
        "disc_C": round(disc.get("C", 0), 1),
        "primary_disc": (personality_type or "S")[0],
        "big5_openness": round(big5.get("openness", 0), 2),
        "big5_conscientiousness": round(big5.get("conscientiousness", 0), 2),
        "big5_extraversion": round(big5.get("extraversion", 0), 2),
        "big5_agreeableness": round(big5.get("agreeableness", 0), 2),
        "big5_neuroticism": round(big5.get("neuroticism", 0), 2),
        "communication_style": communication.get("primary_style"),
        "communication_preferences": "; ".join(communication.get("preferences") or [])
        if isinstance(communication.get("preferences"), list)
        else communication.get("preferences"),
        "sales_approach": sales.get("primary_approach"),
        "decision_style": sales.get("decision_making_style"),
        "motivators": "; ".join(sales.get("key_motivators") or [])
        if isinstance(sales.get("key_motivators"), list)
        else sales.get("key_motivators"),
        "likely_objections": "; ".join(sales.get("likely_objections") or [])
        if isinstance(sales.get("likely_objections"), list)
        else sales.get("likely_objections"),
        "strengths": "; ".join(strengths) if isinstance(strengths, list) else strengths,
        "work_style": (behavioral or {}).get("work_style"),
        "confidence_score": round(confidence, 2) if isinstance(confidence, (int, float)) else confidence,
        "pitch_style": pitch["pitch_style"],
        "pitch_opener_angle": pitch["opener_angle"],
        "pitch_do": pitch["do"],
        "pitch_avoid": pitch["avoid"],
        "data_note": "Internal LeadLab psychometrics (job title + sector + LinkedIn URL; no live LinkedIn scrape)",
    }


def main() -> int:
    parsed = Path(r"c:\Users\attia\Downloads\SERENIDIEN_Prospects_parsed.json")
    leads = json.loads(parsed.read_text(encoding="utf-8"))
    analyzer = InternalPsychometricAnalyzer()

    rows = []
    for i, lead in enumerate(leads, 1):
        rows.append(analyze_one(analyzer, lead, i))
        if i % 50 == 0:
            print(f"Analyzed {i}/{len(leads)}")

    out = Path(r"c:\Users\attia\Downloads\SERENIDIEN_Prospects_Psychometrics.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Psychometrics + Pitch"

    keys = list(rows[0].keys()) if rows else []
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    ws.append(keys)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for r in rows:
        ws.append([r.get(k) for k in keys])

    for col in range(1, len(keys) + 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = min(28, max(12, len(keys[col - 1]) + 2))

    # Summary sheet
    summary = wb.create_sheet("Profile Summary", 0)
    from collections import Counter

    c = Counter((r.get("primary_disc") or "?") for r in rows)
    summary.append(["Primary DISC", "Count", "Share %", "Pitch style"])
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = header_font
    total = len(rows) or 1
    for disc, n in sorted(c.items(), key=lambda x: -x[1]):
        pitch = pitch_for(disc)
        summary.append([disc, n, round(100 * n / total, 1), pitch["pitch_style"]])
    summary.append([])
    summary.append(["Total leads", len(rows)])
    summary.append(["With email", sum(1 for r in rows if r.get("email"))])
    summary.append(["With LinkedIn", sum(1 for r in rows if r.get("linkedin"))])
    summary.append(
        [
            "Note",
            "Profiles from LeadLab internal analyzer. Upload to production to persist on each lead record.",
        ]
    )

    # Pitch playbook sheet
    play = wb.create_sheet("Pitch Playbook by DISC")
    play.append(["DISC", "Pitch style", "Opener angle", "Do", "Avoid"])
    for cell in play[1]:
        cell.fill = header_fill
        cell.font = header_font
    for k, v in PITCH_BY_TYPE.items():
        play.append([k, v["pitch_style"], v["opener_angle"], v["do"], v["avoid"]])
    for col in range(1, 6):
        play.column_dimensions[get_column_letter(col)].width = 40

    wb.save(out)
    print(f"Wrote {out}")
    print("DISC mix:", dict(c))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
